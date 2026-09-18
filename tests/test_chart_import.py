"""Testes do Subsistema de Importação Multiformato de Cifras (v0.4).

Verifica rigorosamente:
1. TextChartSource: texto direto e arquivo .txt.
2. DocxChartSource: documentos .docx com seções, parágrafos e tabelas.
3. ImageChartExtractor:
   - Reconstrução espacial 2D de acordes posicionados sobre a letra
   - Identificação e preservação de tokens de baixa confiança (< 0.70)
   - Zero modificações silenciosas de acordes
4. Convergência: todas as fontes (TXT, DOCX, Imagem, texto bruto) convergem para
   o modelo interno canônico (RawChartDocument -> ChartParser -> ChordNormalizer -> ChordChart).
5. ProjectManager.import_from_source para múltiplos formatos com integração ao Setlist ativo.
"""

import os
import tempfile
import unittest
from PIL import Image, ImageDraw

from app.input.chart_sources import (
    ExtractedToken,
    RawChartDocument,
    TextChartSource,
    DocxChartSource,
    ImageChartSource,
    PdfChartSource
)
from app.input.image_extractor import ImageChartExtractor
from app.input.chart_parser import ChartParser
from app.music.chord_chart import ChordChart
from app.project.project_manager import ProjectManager


class TestChartImport(unittest.TestCase):
    """Testes de importação multiformato de cifras musicais."""

    def setUp(self):
        self.pm = ProjectManager()

    # -------------------------------------------------------------
    # 1. TextChartSource
    # -------------------------------------------------------------
    def test_text_chart_source_from_string(self):
        """Valida carregamento de texto puro colado diretamente."""
        raw_text = (
            "Tom: G\n"
            "BPM: 120\n\n"
            "[Intro]\n"
            "G  D  Em  C\n\n"
            "[Verso 1]\n"
            "G             D\n"
            "Luz do sol que brilha forte\n"
            "Em            C\n"
            "Pelo céu de norte a sul\n"
        )
        source = TextChartSource()
        doc = source.load(raw_text)

        self.assertIsInstance(doc, RawChartDocument)
        self.assertEqual(doc.source_type, "text")
        self.assertEqual(doc.overall_confidence, 1.0)
        self.assertFalse(doc.has_low_confidence_chords)
        self.assertIn("[Intro]", doc.text)
        self.assertIn("G  D  Em  C", doc.text)

        # Convergência para ChordChart
        chart = ChartParser.parse(doc.text)
        self.assertIsInstance(chart, ChordChart)
        self.assertEqual(len(chart.sections), 2)
        self.assertEqual(chart.sections[0].name, "Intro")
        self.assertEqual(chart.sections[1].name, "Verso 1")

    def test_text_chart_source_from_file(self):
        """Valida carregamento a partir de arquivo .txt em disco."""
        content = "[Refrão]\nC9    G/B    Am7    F\nCantando juntos a canção\n"
        with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False, encoding="utf-8") as f:
            f.write(content)
            temp_path = f.name

        try:
            source = TextChartSource()
            doc = source.load(temp_path)

            self.assertEqual(doc.source_type, "text")
            self.assertEqual(os.path.abspath(doc.source_path), os.path.abspath(temp_path))
            self.assertIn("Cantando juntos a canção", doc.text)
            self.assertEqual(doc.metadata.get("lines_count"), 3)

            chart = ChartParser.parse(doc.text)
            self.assertEqual(chart.sections[0].name, "Refrão")
            chord_symbols = [c.symbol.original_symbol for c in chart.sections[0].chords]
            self.assertEqual(chord_symbols, ["C9", "G/B", "Am7", "F"])
        finally:
            if os.path.exists(temp_path):
                os.unlink(temp_path)

    # -------------------------------------------------------------
    # 2. DocxChartSource
    # -------------------------------------------------------------
    def test_docx_chart_source(self):
        """Valida extração estruturada de documentos .docx."""
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx não instalado no ambiente.")

        doc_file = docx.Document()
        doc_file.add_heading("Hino de Teste", level=1)
        doc_file.add_paragraph("[Intro]")
        doc_file.add_paragraph("A  E/G#  F#m  D")
        doc_file.add_paragraph("[Verso]")
        doc_file.add_paragraph("A         E/G#")
        doc_file.add_paragraph("Tudo se renova outra vez")

        with tempfile.NamedTemporaryFile("w", suffix=".docx", delete=False) as f:
            temp_docx_path = f.name

        try:
            doc_file.save(temp_docx_path)

            source = DocxChartSource()
            raw_doc = source.load(temp_docx_path)

            self.assertEqual(raw_doc.source_type, "docx")
            self.assertIn("[Intro]", raw_doc.text)
            self.assertIn("A  E/G#  F#m  D", raw_doc.text)
            self.assertIn("Tudo se renova outra vez", raw_doc.text)
            self.assertGreaterEqual(raw_doc.metadata.get("paragraphs_count", 0), 4)

            # Convergência para ChordChart
            chart = ChartParser.parse(raw_doc.text)
            self.assertEqual(len(chart.sections), 2)
            self.assertEqual(chart.sections[0].name, "Intro")
            self.assertEqual(chart.sections[1].name, "Verso")
            self.assertIn("F#m", [c.symbol.original_symbol for c in chart.sections[0].chords])
        finally:
            if os.path.exists(temp_docx_path):
                os.unlink(temp_docx_path)

    # -------------------------------------------------------------
    # 3. ImageChartExtractor - Reconstrução Espacial 2D
    # -------------------------------------------------------------
    def test_image_chart_spatial_reconstruction(self):
        """Valida o alinhamento bidimensional: acordes projetados com precisão sobre as palavras da letra."""
        extractor = ImageChartExtractor()

        # Tokens simulados:
        # Linha 1 (y=40): Acordes "C9" em x=20, "G/B" em x=240
        # Linha 2 (y=80): Letra "Eu estava pensando em você"
        #   "Eu" x=20 w=30, "estava" x=60 w=70, "pensando" x=140 w=80, "em" x=230 w=30, "você" x=270 w=50
        tokens = [
            ExtractedToken(text="C9", x=20.0, y=40.0, width=40.0, height=20.0, confidence=0.95, is_chord=True),
            ExtractedToken(text="G/B", x=240.0, y=40.0, width=50.0, height=20.0, confidence=0.95, is_chord=True),
            ExtractedToken(text="Eu", x=20.0, y=80.0, width=30.0, height=20.0, confidence=0.95, is_chord=False),
            ExtractedToken(text="estava", x=60.0, y=80.0, width=70.0, height=20.0, confidence=0.95, is_chord=False),
            ExtractedToken(text="pensando", x=140.0, y=80.0, width=80.0, height=20.0, confidence=0.95, is_chord=False),
            ExtractedToken(text="em", x=230.0, y=80.0, width=30.0, height=20.0, confidence=0.95, is_chord=False),
            ExtractedToken(text="você", x=270.0, y=80.0, width=50.0, height=20.0, confidence=0.95, is_chord=False),
        ]

        result_text = extractor._reconstruct_spatial_chart(tokens, image_width=800)
        lines = result_text.splitlines()

        self.assertEqual(len(lines), 2)
        chord_line = lines[0]
        lyric_line = lines[1]

        # Verifica alinhamento: "C9" deve estar alinhado no início
        self.assertTrue(chord_line.startswith("C9"))
        # Verifica que "G/B" foi posicionado mais à direita, acima de "em você"
        gb_pos = chord_line.find("G/B")
        em_pos = lyric_line.find("em")
        self.assertGreater(gb_pos, 10, "G/B deve estar espaçado à direita sobre a letra")
        self.assertAlmostEqual(gb_pos, em_pos, delta=6)

    # -------------------------------------------------------------
    # 4. ImageChartExtractor - Rastreamento de Confiança e Não-Modificação
    # -------------------------------------------------------------
    def test_image_chart_extractor_confidence_tracking(self):
        """Garante que acordes com baixa confiança (< 0.70) são identificados sem alteração silenciosa."""
        extractor = ImageChartExtractor(confidence_threshold=0.70)

        img = Image.new("RGB", (400, 200), color=(255, 255, 255))

        mock_tokens = [
            ExtractedToken(text="C9", x=10.0, y=20.0, width=30.0, height=20.0, confidence=0.95),
            ExtractedToken(text="Bm7", x=80.0, y=20.0, width=60.0, height=20.0, confidence=0.52),
            ExtractedToken(text="Am7", x=160.0, y=20.0, width=40.0, height=20.0, confidence=0.88),
            ExtractedToken(text="Meu", x=10.0, y=60.0, width=30.0, height=20.0, confidence=0.95),
            ExtractedToken(text="coração", x=50.0, y=60.0, width=60.0, height=20.0, confidence=0.90),
        ]

        extractor._run_ocr = lambda path, pil_img: mock_tokens

        doc = extractor.extract(img)

        self.assertIsInstance(doc, RawChartDocument)
        self.assertEqual(doc.source_type, "image")
        self.assertTrue(doc.has_low_confidence_chords)
        self.assertEqual(len(doc.low_confidence_tokens), 1)
        self.assertEqual(doc.low_confidence_tokens[0].text, "Bm7")
        self.assertLess(doc.low_confidence_tokens[0].confidence, 0.70)
        self.assertTrue(doc.low_confidence_tokens[0].is_low_confidence)

    # -------------------------------------------------------------
    # 5. Convergência Global via ProjectManager.import_from_source
    # -------------------------------------------------------------
    def test_project_manager_import_from_text(self):
        """Valida fluxo completo de import_from_source para texto com convergência a ChordChart."""
        raw_text = (
            "[Intro]\n"
            "D  A/C#  Bm7  G\n"
            "[Refrão]\n"
            "D           A/C#\n"
            "Glória nas alturas\n"
        )
        song, doc = self.pm.import_from_source(
            source_input=raw_text,
            title="Canção da Glória",
            artist="Banda Local",
            key="D Major",
            bpm=110.0
        )

        self.assertEqual(song.title, "Canção da Glória")
        self.assertEqual(song.metadata.get("chart_source_type"), "text")
        self.assertIsNotNone(song.chart_data)
        self.assertEqual(len(song.chart_data["sections"]), 2)
        self.assertEqual(song.chart_data["sections"][0]["name"], "Intro")

        active_setlist = self.pm.get_project().get_active_setlist()
        self.assertIn(song, active_setlist.songs)

    def test_project_manager_import_from_txt_file(self):
        """Valida import_from_source recebendo caminho de arquivo .txt."""
        content = "[Verso]\nE  B  C#m  A\nCaminhando pela estrada\n"
        with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False, encoding="utf-8") as f:
            f.write(content)
            temp_path = f.name

        try:
            song, doc = self.pm.import_from_source(temp_path)
            self.assertEqual(song.metadata.get("chart_source_type"), "text")
            self.assertIsNotNone(song.chart_data)
            self.assertEqual(song.chart_data["sections"][0]["chords"][0]["symbol"]["original_symbol"], "E")
        finally:
            if os.path.exists(temp_path):
                os.unlink(temp_path)

    def test_project_manager_import_from_raw_chart_document(self):
        """Valida import_from_source aceitando diretamente um RawChartDocument intermediário."""
        doc = RawChartDocument(
            text="[Intro]\nF#m  D  A  E\n",
            source_type="image",
            overall_confidence=0.88,
            metadata={"source": "scan_test.png"}
        )
        song, ret_doc = self.pm.import_from_source(doc, title="Escaneada")
        self.assertEqual(song.title, "Escaneada")
        self.assertEqual(song.metadata.get("chart_source_type"), "image")
        self.assertEqual(song.metadata.get("chart_overall_confidence"), 0.88)
        self.assertEqual(song.chart_data["sections"][0]["chords"][0]["symbol"]["original_symbol"], "F#m")

    # -------------------------------------------------------------
    # 6. PdfChartSource - Documentos em PDF
    # -------------------------------------------------------------
    def test_pdf_chart_source_digital(self):
        """Valida extração estruturada de PDF vetorial/digital com layout 2D."""
        try:
            from reportlab.pdfgen import canvas
        except ImportError:
            self.skipTest("reportlab não instalado.")

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
            pdf_path = f.name

        try:
            c = canvas.Canvas(pdf_path)
            c.drawString(50, 800, "Tom: G")
            c.drawString(50, 780, "BPM: 125")
            c.drawString(50, 750, "[Intro]")
            c.drawString(50, 730, "G    D    Em    C")
            c.drawString(50, 700, "[Verso]")
            c.drawString(50, 680, "G             D")
            c.drawString(50, 660, "Caminhando pela estrada afora")
            c.save()

            source = PdfChartSource()
            doc = source.load(pdf_path)

            self.assertIsInstance(doc, RawChartDocument)
            self.assertEqual(doc.source_type, "pdf")
            self.assertEqual(doc.metadata.get("pages_count"), 1)
            self.assertFalse(doc.metadata.get("is_scanned"))
            self.assertIn("G    D    Em    C", doc.text)
            self.assertIn("Caminhando pela estrada", doc.text)

            # Convergência para ChordChart
            chart = ChartParser.parse(doc.text)
            self.assertEqual(chart.key, "G Major")
            self.assertEqual(chart.bpm, 125.0)
            self.assertEqual(len(chart.sections), 2)
            self.assertEqual(chart.sections[0].name, "Intro")
            chord_syms = [ch.symbol.original_symbol for ch in chart.sections[0].chords]
            self.assertEqual(chord_syms, ["G", "D", "Em", "C"])
        finally:
            if os.path.exists(pdf_path):
                os.unlink(pdf_path)

    def test_pdf_chart_source_scanned_with_image(self):
        """Valida fallback de PDF escaneado (sem camada de texto, apenas imagem)."""
        try:
            from reportlab.pdfgen import canvas
        except ImportError:
            self.skipTest("reportlab não instalado.")

        # Cria imagem temporária para embutir no PDF
        img = Image.new("RGB", (300, 150), color=(255, 255, 255))
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as img_f:
            img_path = img_f.name
            img.save(img_path)

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as pdf_f:
            pdf_path = pdf_f.name

        try:
            c = canvas.Canvas(pdf_path)
            c.drawImage(img_path, 50, 600, width=300, height=150)
            c.save()

            # Mock extractor para simular reconhecimento da página escaneada
            class MockExtractor:
                def extract(self, pil_image):
                    return RawChartDocument(
                        text="[Intro]\nC    G    Am    F\n",
                        tokens=[],
                        overall_confidence=0.85,
                        source_type="image"
                    )

            source = PdfChartSource(extractor=MockExtractor())
            doc = source.load(pdf_path)

            self.assertEqual(doc.source_type, "pdf")
            self.assertTrue(doc.metadata.get("is_scanned"))
            self.assertIn("C    G    Am    F", doc.text)
        finally:
            if os.path.exists(img_path):
                os.unlink(img_path)
            if os.path.exists(pdf_path):
                os.unlink(pdf_path)

    def test_project_manager_import_from_pdf_file(self):
        """Valida import_from_source do ProjectManager recebendo um arquivo .pdf."""
        try:
            from reportlab.pdfgen import canvas
        except ImportError:
            self.skipTest("reportlab não instalado.")

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
            pdf_path = f.name

        try:
            c = canvas.Canvas(pdf_path)
            c.drawString(50, 780, "[Refrão]")
            c.drawString(50, 760, "A    E    F#m    D")
            c.drawString(50, 730, "Louvemos todos juntos")
            c.save()

            song, doc = self.pm.import_from_source(pdf_path, title="Hino PDF")
            self.assertEqual(song.title, "Hino PDF")
            self.assertEqual(song.metadata.get("chart_source_type"), "pdf")
            self.assertIsNotNone(song.chart_data)
            self.assertEqual(song.chart_data["sections"][0]["name"], "Refrão")
            chord_syms = [ch["symbol"]["original_symbol"] for ch in song.chart_data["sections"][0]["chords"]]
            self.assertEqual(chord_syms, ["A", "E", "F#m", "D"])
        finally:
            if os.path.exists(pdf_path):
                os.unlink(pdf_path)

    # -------------------------------------------------------------
    # 7. Regressão: acordes que coincidem com palavras comuns (A, E, B)
    # -------------------------------------------------------------
    def test_chord_line_preserves_common_word_chords(self):
        """Acordes A/E/B não podem ser descartados como stopwords em uma linha de acordes."""
        from app.input.chart_semantic_classifier import (
            ChartSemanticClassifier,
            SemanticLineType,
        )

        # Linha só de acordes que também são palavras comuns em PT/EN
        classified = ChartSemanticClassifier.classify_line("E  B  C#m  A", 1)
        self.assertEqual(classified.line_type, SemanticLineType.CHORD)
        symbols = [tok for tok, _pos in classified.chord_tokens]
        self.assertEqual(symbols, ["E", "B", "C#m", "A"])

        # Progressão mínima de dois acordes-stopword ainda é linha de acordes
        only_words = ChartSemanticClassifier.classify_line("A  E", 2)
        self.assertEqual(only_words.line_type, SemanticLineType.CHORD)
        self.assertEqual([t for t, _ in only_words.chord_tokens], ["A", "E"])

        # Letra genuína começando com 'E' (conjunção) NÃO vira linha de acordes
        lyric = ChartSemanticClassifier.classify_line("E o amor venceu tudo", 3)
        self.assertEqual(lyric.line_type, SemanticLineType.LYRIC)


if __name__ == "__main__":
    unittest.main()
